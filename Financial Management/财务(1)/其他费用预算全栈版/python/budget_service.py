"""其他费用预算一次性运行版：Excel 导入、校验、会话审核与导出。

所有数据仅在当前 Python 进程内保存。关闭程序后会自动清空，不创建数据库、不复制附件。
"""
from __future__ import annotations

from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parent.parent
TEMPLATE_PATH = ROOT / "其他费用预算导入模板.xlsx"
THRESHOLDS = {"差旅": Decimal("20"), "推广": Decimal("30"), "咨询": Decimal("30"), "办公": Decimal("15")}
_session_records: list[dict[str, Any]] = []
_next_id = 1


class BudgetError(Exception):
    pass


def _money(value: Any) -> Decimal:
    try:
        result = Decimal(str(value).replace(",", "").replace("¥", "").strip())
    except (InvalidOperation, AttributeError):
        raise BudgetError(f"金额“{value}”不是有效数字。")
    if result < 0:
        raise BudgetError("预算金额不能为负数。")
    return result.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def _threshold(item: str) -> Decimal:
    return next((value for keyword, value in THRESHOLDS.items() if keyword in item), Decimal("20"))


def create_template(path: Path = TEMPLATE_PATH) -> Path:
    wb = Workbook(); ws = wb.active; ws.title = "其他费用预算明细"
    ws.append(["部门", "费用项目", "上年实际金额（万元）", "本年预算金额（万元）", "增长原因及测算依据"])
    for cell in ws[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="17365D")
    for index, width in enumerate([16, 22, 22, 22, 52], 1): ws.column_dimensions[chr(64 + index)].width = width
    ws.freeze_panes = "A2"; ws.sheet_view.showGridLines = False; wb.save(path)
    return path


def _header_key(value: Any) -> str:
    return str(value or "").replace(" ", "").replace("\n", "").strip()


def _assess(last_amount: Decimal, current_amount: Decimal, item: str, reason: str) -> tuple[Decimal | None, str, str, str]:
    threshold = _threshold(item)
    if last_amount == 0 and current_amount > 0:
        if not reason:
            raise BudgetError("新增费用必须填写增长原因及测算依据。")
        return None, "高风险", "新增费用：须核验立项、合同或报价依据。", "请上传立项文件、合同或报价单，并说明新增事项、金额测算口径和预算必要性。"
    growth = Decimal("0") if last_amount == 0 else ((current_amount - last_amount) / last_amount * 100).quantize(Decimal("0.01"))
    if growth > threshold:
        if not reason:
            raise BudgetError("同比超阈值必须填写增长原因及测算依据。")
        return growth, "高风险", f"同比增长 {growth}%，超过 {threshold}% 预警阈值。", f"请核对本年预算金额；如金额无误，请补充增长原因、测算口径及合同/报价/活动计划等附件。"
    if growth < 0:
        return growth, "正常", f"预算下降 {abs(growth)}%，未触发增长预警。", "无需修改；如为主动压降，请保留测算依据。如非预期，请核对上年实际与本年预算金额。"
    return growth, "正常", f"同比增长 {growth}%，未超过 {threshold}% 预警阈值。", "无需修改；请确认金额和预算依据准确后提交。"


def _read_headers(path: Path, label: str) -> list[str]:
    if not path.exists() or path.suffix.lower() != ".xlsx":
        raise BudgetError(f"请选择存在的 .xlsx {label}。")
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
        first_row = next(workbook.active.iter_rows(values_only=True))
    except StopIteration:
        raise BudgetError(f"{label}为空，第一行应包含字段名称。")
    except Exception as exc:
        raise BudgetError(f"无法读取{label}：{exc}")
    return [_header_key(value) for value in first_row if _header_key(value)]


def parse_budget_excel(source: Path, rule_template: Path | None = None) -> tuple[list[dict[str, Any]], list[str]]:
    if not source.exists() or source.suffix.lower() != ".xlsx": raise BudgetError("请选择存在的 .xlsx 预算明细文件。")
    try: wb = load_workbook(source, data_only=True, read_only=True)
    except Exception as exc: raise BudgetError(f"无法读取 Excel：{exc}")
    ws = wb.active
    try: headers = {_header_key(value): idx for idx, value in enumerate(next(ws.iter_rows(values_only=True)), 1)}
    except StopIteration: raise BudgetError("Excel 为空，请使用系统模板。")
    if rule_template:
        expected_headers = _read_headers(rule_template, "规则模板")
        missing_headers = [header for header in expected_headers if header not in headers]
        if missing_headers:
            details = "、".join(missing_headers)
            return [], [f"字段异常：部门预算文件缺少规则模板要求的字段：{details}。修改建议：请要求部门补齐这些列，并保持表头名称与规则模板第一行一致后重新导入。"]
    aliases = {"department": ["部门", "所属部门"], "cost_item": ["费用项目", "费用科目", "费用明细"], "last_amount": ["上年实际金额（万元）", "上年实际金额", "上年金额（万元）"], "current_amount": ["本年预算金额（万元）", "本年预算金额", "本年预算（万元）"], "reason": ["增长原因及测算依据", "增长原因", "预算依据"]}
    columns = {}
    for field, names in aliases.items():
        found = next((headers[name] for name in names if name in headers), None)
        if found is None: raise BudgetError(f"Excel 缺少必填列：{names[0]}。请下载并使用系统模板。")
        columns[field] = found - 1
    rows, errors = [], []
    for excel_row, values in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not any(value not in (None, "") for value in values): continue
        try:
            department, item, reason = str(values[columns['department']] or "").strip(), str(values[columns['cost_item']] or "").strip(), str(values[columns['reason']] or "").strip()
            if not department or not item: raise BudgetError("部门和费用项目不能为空。")
            last_amount, current_amount = _money(values[columns['last_amount']]), _money(values[columns['current_amount']])
            growth, risk, message, suggestion = _assess(last_amount, current_amount, item, reason)
            rows.append({"department": department, "cost_item": item, "last_amount": last_amount, "current_amount": current_amount, "growth_rate": growth, "threshold": _threshold(item), "risk_level": risk, "validation_message": message, "suggestion": suggestion, "reason": reason})
        except (BudgetError, IndexError) as exc:
            errors.append(f"第 {excel_row} 行：{exc} 修改建议：请按系统模板补齐字段并修正金额格式后重新导入。")
    if not rows and not errors: errors.append("Excel 未包含任何预算数据。")
    return rows, errors


def submit_budget_excel(source: Path, budget_year: str, attachment: Path | None = None, rule_template: Path | None = None) -> tuple[int, int]:
    global _next_id
    rows, errors = parse_budget_excel(source, rule_template)
    if errors: raise BudgetError("导入被阻断：\n" + "\n".join(errors))
    duplicates = [(row['department'], row['cost_item']) for row in rows if any(existing['budget_year'] == budget_year and existing['department'] == row['department'] and existing['cost_item'] == row['cost_item'] and existing['status'] == '待审核' for existing in _session_records)]
    if duplicates:
        joined = "；".join(f"{department}/{item}" for department, item in duplicates)
        raise BudgetError(f"当前会话中已有待审核记录：{joined}。请先在审核队列处理，或重启程序后重新开始本次测算。")
    for row in rows:
        row.update({"id": _next_id, "budget_year": budget_year, "status": "待审核", "source_file": source.name, "attachment_name": attachment.name if attachment and attachment.exists() else "", "review_note": ""})
        _session_records.append(row); _next_id += 1
    return len(rows), len(rows)


def list_budgets(status: str | None = None) -> list[dict[str, Any]]:
    return [row.copy() for row in _session_records if status is None or row['status'] == status]


def review_budget(budget_id: int, status: str, note: str) -> None:
    if status not in {"已通过", "已退回"}: raise BudgetError("审核状态无效。")
    record = next((row for row in _session_records if row['id'] == budget_id), None)
    if not record or record['status'] != "待审核": raise BudgetError("该记录不存在或已被处理。")
    record['status'], record['review_note'] = status, note.strip()


def export_summary(output: Path) -> tuple[int, Decimal]:
    """导出当次会话的完整明细及按部门审核状态汇总。"""
    rows = list_budgets()
    if not rows: raise BudgetError("当前没有已提交预算，无法导出汇总表。")
    wb = Workbook(); detail = wb.active; detail.title = "预算审核明细"
    detail.append(["预算年度", "部门", "费用项目", "上年实际（万元）", "本年预算（万元）", "增减额（万元）", "同比增长", "预警阈值", "风险标记", "审核状态", "财务意见", "修改建议"])
    for row in rows:
        last, current = row['last_amount'], row['current_amount']
        detail.append([row['budget_year'], row['department'], row['cost_item'], float(last), float(current), float(current-last), float(row['growth_rate'])/100 if row['growth_rate'] is not None else "新增", float(row['threshold'])/100, row['risk_level'], row['status'], row['review_note'], row['suggestion']])
    summary = wb.create_sheet("按部门汇总")
    summary.append(["预算年度", "部门", "已通过金额（万元）", "已退回金额（万元）", "待审核金额（万元）", "合计申报金额（万元）", "已通过项目数", "已退回项目数", "待审核项目数"])
    grouped: dict[tuple[str, str], dict[str, Decimal | int]] = {}
    for row in rows:
        key = (row['budget_year'], row['department'])
        item = grouped.setdefault(key, {"已通过": Decimal("0"), "已退回": Decimal("0"), "待审核": Decimal("0"), "已通过数": 0, "已退回数": 0, "待审核数": 0})
        item[row['status']] += row['current_amount']; item[row['status'] + "数"] += 1
    total = Decimal("0")
    for (year, department), values in sorted(grouped.items()):
        approved, returned, pending = values['已通过'], values['已退回'], values['待审核']
        total += approved
        summary.append([year, department, float(approved), float(returned), float(pending), float(approved + returned + pending), values['已通过数'], values['已退回数'], values['待审核数']])
    for sheet in (detail, summary):
        for cell in sheet[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="17365D")
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions; sheet.sheet_view.showGridLines = False
    for col in [4, 5, 6]:
        for cell in list(detail.columns)[col-1][1:]: cell.number_format = '#,##0.00'
    for col in [7, 8]:
        for cell in list(detail.columns)[col-1][1:]:
            if isinstance(cell.value, (float, int)): cell.number_format = '0.0%'
    for col in [3, 4, 5, 6]:
        for cell in list(summary.columns)[col-1][1:]: cell.number_format = '#,##0.00'
    for i, width in enumerate([13,16,22,18,18,18,14,14,12,12,28,58], 1): detail.column_dimensions[chr(64+i)].width = width
    for i, width in enumerate([13,16,20,20,20,20,16,16,16], 1): summary.column_dimensions[chr(64+i)].width = width
    output.parent.mkdir(parents=True, exist_ok=True); wb.save(output)
    return len(rows), total
