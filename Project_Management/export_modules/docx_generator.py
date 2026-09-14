#!/usr/bin/env python3
"""DOCX export for PMO project monthly report.

Generates a Word document with the same 6-section structure as the HTML/PDF.
Uses python-docx library.

Usage:
    from export_modules.docx_generator import export_docx
    export_docx(snap, output_path)
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any, Optional

from data_modules.metrics_engine import create_metrics, number
from export_modules.html_generator import resolve_page_ids

# ── python-docx import ──
try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
except ImportError:
    Document = None


# ── Color constants ──
BLUE = RGBColor(0x28, 0x64, 0xF3)
NAVY = RGBColor(0x17, 0x3E, 0x79)
INK = RGBColor(0x1C, 0x29, 0x41)
MUTED = RGBColor(0x65, 0x77, 0x93)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
YELLOW = RGBColor(0xD6, 0x8C, 0x00)
RED = RGBColor(0xE3, 0x28, 0x28)
GRAY = RGBColor(0x8B, 0x98, 0xAA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

LABELS = {"normal": "正常", "concern": "关注", "intervention": "干预", "unconfigured": "参考", "unavailable": "待补"}
COLORS = {"normal": GREEN, "concern": YELLOW, "intervention": RED, "unconfigured": GRAY, "unavailable": GRAY}
TARGETS = {
    "req_verify_rate": "≥85%",
    "sprint_completion_rate": "≥85%",
    "workhour_deviation_rate": "±15%",
    "budget_deviation_rate": "±10%",
    "requirement_cost": "≤3人天",
    "defect_escape_rate": "≤3%",
    "bug_close_rate": "≥95%",
    "mttr": "P0≤4h P1≤24h",
    "rd_resource_input_ratio": "研发人力/组织总人数",
}


def n(value, default=0.0):
    parsed = number(value)
    return default if parsed is None else parsed


def metric_value(metric):
    if metric.display_value:
        return metric.display_value
    if metric.value is None:
        return "暂无数据"
    if metric.unit == "%":
        return f"{metric.value:.1%}"
    if metric.unit == "小时":
        return f"{metric.value:.2f} 小时"
    if metric.unit == "个":
        return f"{metric.value:.0f}"
    return f"{metric.value:.2f}"


def state_label(metric):
    return LABELS.get(metric.warning_level, "待补")


def _set_cell_shading(cell, color_hex: str):
    """Set cell background color."""
    shading = cell._element.get_or_add_tcPr()
    shading_elem = shading.makeelement(qn("w:shd"), {
        qn("w:fill"): color_hex,
        qn("w:val"): "clear",
    })
    shading.append(shading_elem)


def _add_styled_paragraph(doc_or_cell, text_str: str, *, size: int = 11,
                          bold: bool = False, color: RGBColor = INK,
                          alignment=None, space_after: int = 6):
    """Add a paragraph with styling."""
    p = doc_or_cell.add_paragraph()
    run = p.add_run(text_str)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if alignment:
        p.alignment = alignment
    p.paragraph_format.space_after = Pt(space_after)
    return p


def _add_metric_table(doc, metrics, proj):
    """Add the 9-metric overview table."""
    table = doc.add_table(rows=len(metrics) + 1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    headers = ["指标名称", "当前值", "目标值", "计算口径", "预警等级"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(9)
                run.font.bold = True
                run.font.color.rgb = WHITE
        _set_cell_shading(cell, "2864F3")

    # Data rows
    for row_idx, m in enumerate(metrics):
        row = table.rows[row_idx + 1]
        values = [
            m.name,
            metric_value(m),
            TARGETS.get(m.code, ""),
            m.formula[:60] + ("..." if len(m.formula) > 60 else ""),
            state_label(m),
        ]
        for col_idx, val in enumerate(values):
            cell = row.cells[col_idx]
            cell.text = str(val)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)
                    run.font.color.rgb = INK
            # Color the warning level cell
            if col_idx == 4:
                level_color = {"正常": "E9F9EE", "关注": "FFF8DC", "干预": "FFEBEB"}.get(val, "EDF1F6")
                _set_cell_shading(cell, level_color)

        # Alternate row shading
        if row_idx % 2 == 1:
            for col_idx in range(4):
                _set_cell_shading(row.cells[col_idx], "F8FAFD")

    return table


def _add_section_header(doc, title: str, question: str, decision: str):
    """Add a section header with title, question, and decision."""
    _add_styled_paragraph(doc, title, size=16, bold=True, color=NAVY, space_after=2)
    _add_styled_paragraph(doc, f"核心问题：{question}", size=10, color=MUTED, space_after=4)
    # Decision box as indented paragraph
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(f"▸ {decision}")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x35, 0x59, 0x8D)


def _add_notes(doc, left_title: str, left_body: str, right_title: str, right_body: str):
    """Add conclusion and action notes."""
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for col, (title, body) in enumerate([(left_title, left_body), (right_title, right_body)]):
        # Title row
        cell = table.rows[0].cells[col]
        cell.text = title
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(10)
                r.font.bold = True
                r.font.color.rgb = NAVY
        _set_cell_shading(cell, "EDF4FF")

        # Body row
        cell = table.rows[1].cells[col]
        cell.text = body
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(9)
                r.font.color.rgb = MUTED

    # Set column widths
    for row in table.rows:
        for cell in row.cells:
            cell.width = Cm(8)


def _add_data_table(doc, title: str, headers: list[str], rows: list[list]):
    """Add a data table with headers and rows."""
    _add_styled_paragraph(doc, title, size=11, bold=True, color=INK, space_after=4)

    if not rows:
        _add_styled_paragraph(doc, "（暂无足够数据展示趋势）", size=9, color=GRAY)
        return

    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.size = Pt(8)
                r.font.bold = True
                r.font.color.rgb = WHITE
        _set_cell_shading(cell, "2864F3")

    # Data
    for row_idx, row_data in enumerate(rows):
        for col_idx, val in enumerate(row_data):
            cell = table.rows[row_idx + 1].cells[col_idx]
            cell.text = str(val) if val is not None else "—"
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
                    r.font.color.rgb = INK


def export_docx(snap: dict, output_path: Path, page_ids=None) -> Path:
    """Generate a Word document from snapshot data.

    Args:
        snap: snapshot dict (from fetch_data.fetch)
        output_path: output .docx file path

    Returns:
        Path to the generated file
    """
    if Document is None:
        raise ImportError("python-docx 未安装: pip install python-docx")

    proj = snap.get("project", {})
    bug_detail = snap.get("bug_detail", {})
    staff_list = snap.get("staff_daily_workhours", [])
    quality_history = snap.get("quality_history", [])

    metrics, _ = create_metrics(snap)
    items = {m.code: m for m in metrics}

    closed_bug_count = bug_detail.get("bug_closed_count") or 0
    total_bug_count = bug_detail.get("bug_total_count") or 0

    project_name = proj.get("project_name", "未知项目")
    month = proj.get("stat_month", "")
    manager = proj.get("manager_name", "待补")
    project_code = proj.get("project_code", proj.get("project_id", ""))
    product_line = proj.get("product_line") or proj.get("department_name") or "产品线"

    # ── Create document ──
    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    # ── Title page ──
    _add_styled_paragraph(doc, "项目管理快报", size=24, bold=True, color=NAVY,
                          alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    _add_styled_paragraph(doc, f"{product_line} · {project_name}", size=14, color=INK,
                          alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    _add_styled_paragraph(doc, f"（{project_code}）", size=10, color=MUTED,
                          alignment=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    # Project info table
    info_table = doc.add_table(rows=3, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("报告期", month),
        ("项目经理", manager),
        ("数据状态", "月度快报"),
    ]
    for i, (label, value) in enumerate(info_data):
        info_table.rows[i].cells[0].text = label
        info_table.rows[i].cells[1].text = str(value)
        for p in info_table.rows[i].cells[0].paragraphs:
            for r in p.runs:
                r.font.size = Pt(10)
                r.font.bold = True
                r.font.color.rgb = NAVY
        for p in info_table.rows[i].cells[1].paragraphs:
            for r in p.runs:
                r.font.size = Pt(10)
                r.font.color.rgb = INK

    # ── Metric overview ──
    doc.add_paragraph()  # spacer
    _add_styled_paragraph(doc, "本期核心指标", size=14, bold=True, color=NAVY, space_after=6)
    _add_metric_table(doc, metrics, proj)

    # ── Section 1: 产研精细化管理 ──
    doc.add_page_break()
    task_done = n(proj.get("task_finish_count"))
    task_total = n(proj.get("task_count"))
    task_delay = n(proj.get("task_delay_count"))
    _add_section_header(doc,
        "① 产研精细化管理", "交付节奏是否可控？",
        f"本期完成 {task_done:.0f}/{task_total:.0f} 项任务，延期 {task_delay:.0f} 项。")
    _add_data_table(doc, "任务完成情况",
        ["类别", "数量"],
        [["任务总数", f"{task_total:.0f}"],
         ["已完成", f"{task_done:.0f}"],
         ["延期任务", f"{task_delay:.0f}"],
         ["高难任务", f"{n(proj.get('high_difficulty_task_count')):.0f}"]])
    _add_notes(doc,
        "本期结论", f"Sprint 完成率 {metric_value(items['sprint_completion_rate'])}，交付结果{'正常' if items['sprint_completion_rate'].warning_level == 'normal' else '需关注'}。",
        "建议动作", "按需求变更、资源不足、技术阻塞拆分延期原因，将完成率与延期率纳入例会复盘。")

    # ── Section 2: 项目投入偏差 ──
    doc.add_page_break()
    actual_hours = n(proj.get("actual_total_man_hour"))
    planned_hours = n(proj.get("plan_total_man_hour"))
    _add_section_header(doc,
        "② 项目投入偏差", "钱和时间花在哪？偏了多少？",
        f"实际投入 {actual_hours:,.0f} 小时，计划 {planned_hours:,.0f} 小时。")
    _add_notes(doc,
        "偏差判断", f"工时偏差率 {metric_value(items['workhour_deviation_rate'])}，预算偏差率 {metric_value(items['budget_deviation_rate'])}。",
        "建议动作", "按新增需求、返工、资源价格、外包/人力拆解成本偏差；确认计划基线是否仍反映当前范围。")

    # ── Section 3: 人员投入健康度 ──
    doc.add_page_break()
    _add_section_header(doc,
        "③ 人员投入健康度", "团队是否在透支？",
        "已有项目成员投入明细，加班与跨项目投入待接入日考勤数据。")
    staff_sorted = sorted(staff_list, key=lambda r: n(r.get("allocated_hour")), reverse=True)[:10]
    if staff_sorted:
        _add_data_table(doc, "项目成员投入工时 Top 10",
            ["成员", "投入工时"],
            [[r.get("staff_name", "成员"), f"{n(r.get('allocated_hour')):.1f}"] for r in staff_sorted])
    _add_notes(doc,
        "已确认口径", "实际总工时是所有成员在该项目的投入总和；9 小时/人天只用于需求成本换算。",
        "待接入数据", "成员逐日项目投入时长、加班工时、跨项目投入明细。")

    # ── Section 4: 产品质量 ──
    doc.add_page_break()
    _add_section_header(doc,
        "④ 产品质量", "交付物靠不靠谱？",
        f"缺陷逃逸率 {metric_value(items['defect_escape_rate'])}；Bug 关闭率 {metric_value(items['bug_close_rate'])}。")
    _add_data_table(doc, "Bug 关闭情况",
        ["状态", "数量"],
        [["已关闭", f"{closed_bug_count}"],
         ["未关闭", f"{max(total_bug_count - closed_bug_count, 0)}"],
         ["总计", f"{total_bug_count}"]])
    _add_notes(doc,
        "当前质量状态", f"缺陷逃逸率 {metric_value(items['defect_escape_rate'])}；当前缺少线上 Bug 数和总 Bug 数复核。",
        "建议动作", "持续跟踪 Bug 关闭率与缺陷逃逸率趋势，确保质量判断链路完整。")

    # ── Section 5: Bug 修复效率 ──
    doc.add_page_break()
    mttr_val = items['mttr'].value or 0
    _add_section_header(doc,
        "⑤ Bug 修复效率", "问题解决得快不快？",
        f"本期关闭 {closed_bug_count}/{total_bug_count} 个 Bug，MTTR {mttr_val:.2f} 小时。")
    _add_notes(doc,
        "计算公式", f"排除拒绝/第三方后 Bug 修复时长 ÷ Bug 总数({total_bug_count}) ÷ 3600 = {mttr_val:.2f} 小时。",
        "后续接入", "接入每个 Bug 的严重级别和状态，以计算分级 MTTR 与月度趋势。")

    # ── Section 6: 需求价值交付 ──
    doc.add_page_break()
    _add_section_header(doc,
        "⑥ 需求价值交付", "做的东西有没有用？",
        f"需求成本 {metric_value(items['requirement_cost'])}，研发资源占比 {metric_value(items['rd_resource_input_ratio'])}。")
    _add_data_table(doc, "工时投入结构",
        ["类型", "工时"],
        [["研发工时", f"{n(proj.get('actual_rd_man_hour')):,.1f}"],
         ["测试工时", f"{n(proj.get('actual_test_man_hour')):,.1f}"]])
    _add_notes(doc,
        "研发资源投入占比", f"项目参与人员 {metric_value(items['rd_resource_input_ratio'])}（{items['rd_resource_input_ratio'].note}）。",
        "待建立价值闭环", "接入需求验收、活跃使用、客户/业务收益数据，建立投入-交付-验证-收益链路。")

    # ── Footer ──
    doc.add_paragraph()
    _add_styled_paragraph(doc,
        f"项目管理快报 · 由 PMO 项目月度快报工作流生成 · {month}",
        size=8, color=GRAY, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    # ── Save ──
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path
